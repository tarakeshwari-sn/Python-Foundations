import os
import shutil
import csv
import logging
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor
from openpyxl import Workbook,load_workbook

incoming_dir="incoming_files"
processed_dir="processed_data"
destinations={".csv": "sales",".json": "logs",".txt": "errors",".xlsx":"reports"}
log_file="automation.log"
MAX_WORKERS=5

class FileWorker:
    def __init__(self, file_path):
        self.file_path=file_path
        self.extension=os.path.splitext(file_path)[1].lower()

    def process(self):
        if self.extension not in destinations:
            return

        dest_folder = destinations[self.extension]
        target_dir = os.path.join(processed_dir, dest_folder)
        os.makedirs(target_dir, exist_ok=True)

        if self.extension == ".csv":
            self._process_csv()
        elif self.extension == ".xlsx":
            self._process_excel()

        self._move_file(target_dir)


    def _process_csv(self):
        total=0.0
        try:
            with open(self.file_path,newline="",encoding="utf-8") as f:
                reader=csv.DictReader(f)
                for row in reader:
                    total+=float(row.get("Amount", 0))
            
            logging.info(f"Processed CSV {os.path.basename(self.file_path)} | Total Amount: {total}")
        
        except Exception as e:
            logging.error(f"Failed processing CSV {self.file_path} | Error: {e}")

    def _process_excel(self):
        total=0.0
        try:
            wb=load_workbook(self.file_path, read_only=True, data_only=True)
            ws=wb.active
            header=[cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            if "Score" not in header:
                logging.warning(f"Excel {os.path.basename(self.file_path)} missing 'Score' column")
                return

            score_index = header.index("Score")
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[score_index] is not None:
                    total += float(row[score_index])

            wb.close()
            logging.info(f"Processed EXCEL {os.path.basename(self.file_path)} | Total Score: {total}")

        except Exception as e:
            logging.error(f"Failed processing EXCEL {self.file_path} | Error: {e}")

    def _move_file(self, target_dir):
        try:
            shutil.move(self.file_path, target_dir)
            logging.info(f"Moved {os.path.basename(self.file_path)} to {target_dir}")
        except Exception as e:
            logging.error(f"Failed moving file {self.file_path} | Error: {e}")

class FileOrganizer:
    def __init__(self, incoming_dir):
        self.incoming_dir=incoming_dir

    def scan_files(self):
        files=[]
        for file in os.listdir(self.incoming_dir):
            full_path=os.path.join(self.incoming_dir, file)
            if os.path.isfile(full_path):
                files.append(full_path)
        return files

    def run(self):
        files = self.scan_files()
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            for file_path in files:
                worker=FileWorker(file_path)
                executor.submit(worker.process)

def setup_logging():
    logging.basicConfig(filename=log_file,level=logging.INFO,
        format="[%(levelname)s] %(asctime)s - %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S")

if __name__ == "__main__":
    setup_logging()
    organizer=FileOrganizer(incoming_dir)
    organizer.run()
